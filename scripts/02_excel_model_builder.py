import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAW_FILE = "data/raw/urban_cart_retail_dataset.xlsx"
AUDITED_FILE = "data/processed/urban_cart_financial_model.xlsx"

print("Loading raw dataset into openpyxl...")
wb = openpyxl.load_workbook(RAW_FILE)

# --- 1. Append Unit Economic Columns to fact_orders ---
ws_orders = wb["fact_orders"]
max_row = ws_orders.max_row

new_headers = [
    ("net_sales", "R"),
    ("cm1_gross_margin", "S"),
    ("net_logistics_cost", "T"),
    ("cm2_variable_profit", "U"),
    ("cm2_margin_pct", "V"),
    ("is_profitable_order", "W")
]

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

for title, col_letter in new_headers:
    cell = ws_orders[f"{col_letter}1"]
    cell.value = title
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

print(f"Applying unit economics formulas to {max_row - 1} rows...")

# Formula logic:
# I: gross_sales, J: item_cost, K: promo_discount
# L: delivery_fee_charged, M: actual_delivery_cost, N: payment_gateway_fee
for r in range(2, max_row + 1):
    # Net Sales = Gross Sales - Promo Discount
    ws_orders[f"R{r}"] = f"=I{r}-K{r}"
    
    # CM1 (Gross Margin) = Net Sales - Wholesale Item Cost
    ws_orders[f"S{r}"] = f"=R{r}-J{r}"
    
    # Net Logistics Deficit = Customer Delivery Fee - Fleet Cost
    ws_orders[f"T{r}"] = f"=L{r}-M{r}"
    
    # CM2 (Variable Profit) = CM1 + Net Logistics - PG Fee
    ws_orders[f"U{r}"] = f"=S{r}+T{r}-N{r}"
    
    # Margin % (divide by zero guard)
    ws_orders[f"V{r}"] = f'=IF(R{r}>0, U{r}/R{r}, 0)'
    
    # Order profitability flag
    ws_orders[f"W{r}"] = f'=IF(U{r}>0, 1, 0)'
    
    # Accounting number formats
    ws_orders[f"R{r}"].number_format = '#,##0.00'
    ws_orders[f"S{r}"].number_format = '#,##0.00'
    ws_orders[f"T{r}"].number_format = '#,##0.00'
    ws_orders[f"U{r}"].number_format = '#,##0.00'
    ws_orders[f"V{r}"].number_format = '0.0%'
    ws_orders[f"W{r}"].number_format = '0'

# --- 2. Executive Scenario & Sensitivity Sheet ---
print("Building executive scenario sheet...")
if "executive_scenario_model" in wb.sheetnames:
    del wb["executive_scenario_model"]

ws_model = wb.create_sheet(title="executive_scenario_model")
ws_model.views.sheetView[0].showGridLines = True

# Style definitions
title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
section_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
kpi_label_font = Font(name="Calibri", size=11, bold=True)
kpi_val_font = Font(name="Calibri", size=11, bold=False)
highlight_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

# Sheet Title
ws_model["B2"] = "URBANCART: EXECUTIVE UNIT ECONOMICS & SCENARIO ENGINE"
ws_model["B2"].font = title_font

# Section 1: Baseline Portfolio KPIs
ws_model["B4"] = "Baseline Portfolio KPIs (Actual 7,000 Orders)"
ws_model.merge_cells("B4:D4")
ws_model["B4"].font = section_font
ws_model["B4"].fill = section_fill

kpis = [
    ("Total Gross Merchandise Value (GMV)", f"=SUM(fact_orders!I2:I{max_row})", "₹#,##0.00"),
    ("Total Promo Discounts Burned", f"=SUM(fact_orders!K2:K{max_row})", "₹#,##0.00"),
    ("Discount Burn Rate (%)", "=C6/C5", "0.0%"),
    ("Net Sales", f"=SUM(fact_orders!R2:R{max_row})", "₹#,##0.00"),
    ("Total Direct Product Cost (COGS)", f"=SUM(fact_orders!J2:J{max_row})", "₹#,##0.00"),
    ("Gross Contribution Margin 1 (CM1)", f"=SUM(fact_orders!S2:S{max_row})", "₹#,##0.00"),
    ("Net Logistics Deficit (Delivery Burn)", f"=SUM(fact_orders!T2:T{max_row})", "₹#,##0.00"),
    ("Total Payment Gateway Fees", f"=SUM(fact_orders!N2:N{max_row})", "₹#,##0.00"),
    ("Net Contribution Margin 2 (CM2)", f"=SUM(fact_orders!U2:U{max_row})", "₹#,##0.00"),
    ("Portfolio CM2 Margin (%)", "=C13/C8", "0.0%"),
    ("Order Success Rate (Delivered vs Total)", f'=COUNTIF(fact_orders!G2:G{max_row}, "DELIVERED")/COUNTA(fact_orders!G2:G{max_row})', "0.0%"),
    ("Return Leakage Rate", f'=COUNTIF(fact_orders!G2:G{max_row}, "RETURNED")/COUNTA(fact_orders!G2:G{max_row})', "0.0%"),
    ("Profitable Order Proportion", f"=AVERAGE(fact_orders!W2:W{max_row})", "0.0%")
]

for idx, (label, formula, num_fmt) in enumerate(kpis, start=5):
    ws_model[f"B{idx}"] = label
    ws_model[f"B{idx}"].font = kpi_label_font
    ws_model[f"B{idx}"].border = thin_border
    
    ws_model[f"C{idx}"] = formula
    ws_model[f"C{idx}"].font = kpi_val_font
    ws_model[f"C{idx}"].number_format = num_fmt
    ws_model[f"C{idx}"].alignment = Alignment(horizontal="right")
    ws_model[f"C{idx}"].border = thin_border

# Section 2: Sensitivity Table - Free Delivery Threshold
ws_model["F4"] = "Free-Delivery Threshold Sensitivity Simulation"
ws_model.merge_cells("F4:I4")
ws_model["F4"].font = section_font
ws_model["F4"].fill = section_fill

sim_headers = ["Min Basket Size for Free Delivery", "Projected Delivery Fee Recovery", "Net CM2 Realization", "Margin Expansion"]
for col_idx, text in enumerate(sim_headers, start=6):
    cell = ws_model.cell(row=5, column=col_idx)
    cell.value = text
    cell.font = Font(bold=True)
    cell.fill = highlight_fill
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

threshold_scenarios = [
    (299, 1.25),
    (349, 1.15),
    (399, 1.00),  # Baseline
    (449, 0.88),
    (499, 0.76)
]

for row_idx, (threshold, fee_multiplier) in enumerate(threshold_scenarios, start=6):
    ws_model[f"F{row_idx}"] = f"Threshold ₹{threshold}"
    ws_model[f"G{row_idx}"] = f"=C11*{fee_multiplier}"
    ws_model[f"H{row_idx}"] = f"=C13+(G{row_idx}-C11)"
    ws_model[f"I{row_idx}"] = f"=(H{row_idx}-C13)/C13"
    
    ws_model[f"F{row_idx}"].border = thin_border
    ws_model[f"G{row_idx}"].border = thin_border
    ws_model[f"H{row_idx}"].border = thin_border
    ws_model[f"I{row_idx}"].border = thin_border
    
    ws_model[f"G{row_idx}"].number_format = "₹#,##0.00"
    ws_model[f"H{row_idx}"].number_format = "₹#,##0.00"
    ws_model[f"I{row_idx}"].number_format = "+0.0%;-0.0%"

# Column width auto-fit
for ws in [ws_orders, ws_model]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

print(f"Saving audited workbook to {AUDITED_FILE}...")
os.makedirs("data/processed", exist_ok=True)
wb.save(AUDITED_FILE)
print("Excel financial model generated successfully.")